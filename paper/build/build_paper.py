# -*- coding: utf-8 -*-
"""Fill the IEEE A4 conference template with the SpecDrift-Bench paper.

The template's body is a two-column continuous section that ends at the last
guidance paragraph; everything after it belongs to a one-column trailer. New
content is therefore inserted *before* that paragraph, so it lands in the
two-column section.
"""
import copy

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph

import content as C

DOC = docx.Document("template_fixed.docx")
BODY = DOC.element.body
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def clear_runs(par):
    for r in list(par.runs):
        r._element.getparent().remove(r._element)


children = list(BODY.iterchildren())

# ------------------------------------------------------------------ title
title_p = Paragraph(children[0], DOC)
keeper = title_p.runs[0]
keeper.text = C.TITLE
for r in title_p.runs[1:]:
    r._element.getparent().remove(r._element)

# The "*Note: Sub-titles ..." line and the spare blank are not part of a paper.
children[1].getparent().remove(children[1])
children[2].getparent().remove(children[2])

# ------------------------------------------------------------------ authors
author_template = children[4]
anchor_author = children[7]
parent = anchor_author.getparent()


def build_authors():
    """One paragraph holding every author, separated by column breaks.

    This is the template's own idiom, and it keeps all four blocks flush with the
    top of the author section; one paragraph per column does not.
    """
    el = copy.deepcopy(author_template)
    par = Paragraph(el, DOC)
    clear_runs(par)
    for a_i, fields in enumerate(C.AUTHORS):
        if a_i:
            run = par.add_run()
            run._element.append(
                run._element.makeelement(qn("w:br"), {qn("w:type"): "column"})
            )
        for i, line in enumerate(fields):
            run = par.add_run(line)
            run.font.size = Pt(8.5)
            run.italic = i in (1, 2)
            if i < len(fields) - 1:
                run.add_break()
    return el


parent.insert(list(parent).index(anchor_author), build_authors())
for old in children[4:8]:
    old.getparent().remove(old)

# Four authors, one per column of the author section.
for cols in BODY.iter(qn("w:cols")):
    if cols.get(qn("w:num")) == "3":
        cols.set(qn("w:num"), "4")

# ------------------------------------------------- find the two-column anchor
ANCHOR = None
for ch in BODY.iterchildren():
    if ch.tag != qn("w:p"):
        continue
    sect = ch.find(".//" + qn("w:sectPr"))
    if sect is None:
        continue
    cols = sect.find(qn("w:cols"))
    if cols is not None and cols.get(qn("w:num")) == "2":
        ANCHOR = ch
assert ANCHOR is not None, "two-column body section not found"
clear_runs(Paragraph(ANCHOR, DOC))

# ------------------------------------------------------------------ clear body
for ch in list(BODY.iterchildren()):
    if ch is ANCHOR or ch.tag == qn("w:sectPr"):
        continue
    if ch.find(".//" + qn("w:sectPr")) is not None:
        continue
    if ch.tag == qn("w:p") and Paragraph(ch, DOC).style.name in ("paper title", "Author"):
        continue
    ch.getparent().remove(ch)

# ------------------------------------------------------------------ emitters
def emit(element):
    """Move a freshly appended element into the two-column body section."""
    element.getparent().remove(element)
    ANCHOR.addprevious(element)
    return element


def para(style=None):
    p = DOC.add_paragraph(style=style)
    emit(p._element)
    return p


BORDERS = parse_xml(
    '<w:tblBorders xmlns:w="%s">' % W
    + "".join(
        '<w:%s w:val="single" w:sz="4" w:space="0" w:color="auto"/>' % edge
        for edge in ("top", "start", "left", "bottom", "end", "right", "insideH", "insideV")
    )
    + "</w:tblBorders>"
)


def add_table(caption, rows, widths):
    cap = para("table head")
    cap.add_run(caption)
    cap.paragraph_format.keep_with_next = True

    t = DOC.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.autofit = False
    t._tbl.tblPr.append(copy.deepcopy(BORDERS))
    total = parse_xml('<w:tblW xmlns:w="%s" w:w="%d" w:type="dxa"/>'
                      % (W, int(sum(widths) * 1440)))
    t._tbl.tblPr.append(total)
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for r_i, row in enumerate(rows):
        for c_i, text in enumerate(row):
            cell = t.cell(r_i, c_i)
            cell.width = Inches(widths[c_i])
            p = cell.paragraphs[0]
            p.style = DOC.styles["table col head" if r_i == 0 else "table copy"]
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if r_i == 0 or c_i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            )
            p.add_run(text)
            # Keep the caption and every row on one page.
            p.paragraph_format.keep_with_next = r_i < len(rows) - 1
    emit(t._tbl)
    para("Body Text")


def add_figure(path, caption):
    p = para()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(3.25))
    para("figure caption").add_run(" " + caption)


# ------------------------------------------------------------------ abstract
ab = para("Abstract")
r = ab.add_run("Abstract—")
r.bold, r.italic = True, True
ab.add_run(C.ABSTRACT)

kw = para("Keywords")
kw.add_run("Keywords—").italic = True
kw.add_run(C.KEYWORDS)

# ------------------------------------------------------------------ body
for item in C.BODY:
    if item[0] == "TABLE":
        add_table(item[1], item[2], item[3])
    elif item[0] == "FIGURE":
        add_figure(item[1], item[2])
    else:
        para(item[0]).add_run(item[1])

DOC.core_properties.title = C.TITLE
DOC.core_properties.author = ", ".join(a[0] for a in C.AUTHORS)
DOC.save("../SpecDrift-Bench_IEEE_paper.docx")
print("saved")
